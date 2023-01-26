import docx


doc = docx.Document(r'C:\Users\Administrator\Downloads\новый договор.docx')

tabs = doc.tables
indx_adress = tabs[0].cell(
                     0, 0).text.strip().split().index('ИНН')

print(indx_adress)
